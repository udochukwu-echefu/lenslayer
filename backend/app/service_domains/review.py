from __future__ import annotations

from pathlib import Path
from typing import Any
import io

from fastapi import HTTPException
from sqlalchemy import select

from ..models import ContractVersion, User
from .common import json_load, safe_filename


class ReviewServiceMixin:
    def redline_export(
        self,
        organization_id: str,
        contract_id: str,
        user: User,
    ) -> tuple[bytes, str]:
        contract = self.contracts.get_contract(organization_id, contract_id, user)
        review = self.contracts.get_review(organization_id, contract_id, user)
        version = self.session.scalar(select(ContractVersion).where(
            ContractVersion.organization_id == organization_id,
            ContractVersion.contract_id == contract_id,
        ).order_by(ContractVersion.version_number.desc()).limit(1))
        asset = version.document_asset if version else None
        if not asset or asset.status != "available" or Path(asset.original_name).suffix.lower() != ".docx":
            raise HTTPException(
                status_code=409,
                detail="Tracked-change redlining requires a retained DOCX version. Upload a DOCX with document retention enabled.",
            )
        findings = json_load(review.analysis_json, {}).get("risk_assessment", [])
        from ..redline import build_redline
        content, change_count = build_redline(self.object_store.get(asset.storage_key), findings)
        if not change_count:
            raise HTTPException(status_code=409, detail="This review has no suggested replacement language to redline.")
        self._audit(organization_id, user.id, "review.redline_exported", contract_id, {
            "version_number": version.version_number,
            "suggested_change_count": change_count,
        })
        self.session.commit()
        return content, f"{safe_filename(Path(contract.source_name).stem)}-lenslayer-redline.docx"

    def answer_contract_question(
        self,
        organization_id: str,
        contract_id: str,
        user: User,
        question: str,
    ) -> tuple[str, list[dict[str, str]], str]:
        review = self.contracts.get_review(organization_id, contract_id, user)
        source_text = (review.source_text or "").strip()
        if not source_text:
            raise HTTPException(
                status_code=409,
                detail="Contract Q&A requires retained source text. Upload the contract again with source-text retention enabled.",
            )
        contract = self.contracts.get_contract(organization_id, contract_id, user)
        result = self.review_workflow.answer_contract(
            source_text,
            question,
            json_load(contract.review_context_json, {}),
        )
        sources = [
            {
                "label": source.label,
                "location": source.location,
                "excerpt": source.excerpt,
            }
            for source in result.sources
        ]
        if not sources:
            raise HTTPException(status_code=409, detail="The retained source text is empty.")
        self._audit(
            organization_id,
            user.id,
            "review.question_answered",
            contract_id,
            {
                "question": question.strip()[:240],
                "source_count": len(sources),
                "generated_by": result.generated_by,
            },
        )
        self.session.commit()
        return result.answer, sources, result.generated_by


    def contract_export(
        self,
        organization_id: str,
        contract_id: str,
        user: User,
        export_format: str,
    ) -> tuple[bytes, str, str]:
        contract = self.contracts.get_contract(organization_id, contract_id, user)
        review = self.contracts.get_review(organization_id, contract_id, user)
        analysis = json_load(review.analysis_json, {})
        quality = json_load(review.quality_json, {})
        context = json_load(contract.review_context_json, {})
        from ..report_exports import (
            build_csv,
            build_docx_report,
            build_json_report,
            build_markdown_report,
            build_pdf_report,
        )

        safe_stem = safe_filename(Path(contract.source_name).stem)
        if export_format == "pdf":
            content = build_pdf_report(analysis, contract.source_name, context)
            media_type, suffix = "application/pdf", "pdf"
        elif export_format == "docx":
            content = build_docx_report(analysis, contract.source_name, context)
            media_type, suffix = "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
        elif export_format == "csv":
            rows: list[dict[str, Any]] = []
            for category, items in (
                ("risk", analysis.get("risk_assessment", [])),
                ("protection_gap", analysis.get("missing_protections", [])),
                ("obligation", analysis.get("obligations", [])),
                ("deadline", analysis.get("deadlines", [])),
                ("payment", analysis.get("payments", [])),
                ("negotiation", analysis.get("negotiation_priorities", [])),
            ):
                for item in items or []:
                    rows.append({"category": category, **(item if isinstance(item, dict) else {"detail": item})})
            content = build_csv(rows).encode("utf-8")
            media_type, suffix = "text/csv; charset=utf-8", "csv"
        elif export_format == "md":
            content = build_markdown_report(analysis, contract.source_name, context).encode("utf-8")
            media_type, suffix = "text/markdown; charset=utf-8", "md"
        elif export_format == "json":
            content = build_json_report(analysis, context, quality).encode("utf-8")
            media_type, suffix = "application/json; charset=utf-8", "json"
        else:
            raise HTTPException(status_code=422, detail="Choose PDF, DOCX, CSV, Markdown, or JSON.")
        self._audit(
            organization_id,
            user.id,
            "review.exported",
            contract_id,
            {"format": export_format},
        )
        self.session.commit()
        return content, media_type, f"{safe_stem}-lenslayer-review.{suffix}"

    def counsel_handoff(
        self,
        organization_id: str,
        contract_id: str,
        user: User,
    ) -> tuple[bytes, str]:
        contract = self.contracts.get_contract(organization_id, contract_id, user)
        review = self.contracts.get_review(organization_id, contract_id, user)
        from docx import Document

        analysis = json_load(review.analysis_json, {})
        document = Document()
        document.add_heading(f"Counsel handoff: {contract.title}", 0)
        document.add_paragraph("Prepared by LensLayer for qualified professional review. Not legal advice.")
        document.add_heading("Executive summary", 1)
        document.add_paragraph(str(analysis.get("executive_summary") or "No executive summary returned."))
        document.add_heading("Open risks and protection gaps", 1)
        for item in analysis.get("risk_assessment", []):
            document.add_heading(str(item.get("title") or "Clause finding"), 2)
            document.add_paragraph(str(item.get("explanation") or ""))
            document.add_paragraph(f"Evidence: {item.get('citation') or 'Not identified'} | {item.get('quote') or 'No quote returned'}")
        for item in analysis.get("missing_protections", []):
            document.add_paragraph(str(item.get("issue") if isinstance(item, dict) else item), style="List Bullet")
        document.add_heading("Negotiation priorities", 1)
        for item in analysis.get("negotiation_priorities", []):
            document.add_paragraph(str(item.get("title") if isinstance(item, dict) else item), style="List Bullet")
        document.add_heading("Human decisions", 1)
        for decision in self.collaboration.list_contract_decisions(organization_id, contract_id, user):
            document.add_paragraph(
                f"{decision.decision.upper()}: {decision.subject} | {decision.reviewer.display_name} | {decision.rationale}",
                style="List Bullet",
            )
        document.add_heading("Open actions and approvals", 1)
        for task in self.tasks.list_tasks(organization_id, user, contract_id=contract_id):
            if task.status not in {"done", "cancelled"}:
                document.add_paragraph(f"{task.title} ({task.status})", style="List Bullet")
        for approval in self.collaboration.list_approval_requests(organization_id, contract_id, user):
            document.add_paragraph(f"{approval.title} ({approval.status})", style="List Bullet")
        output = io.BytesIO()
        document.save(output)
        self._audit(organization_id, user.id, "counsel_handoff.exported", contract_id)
        self.session.commit()
        return output.getvalue(), f"{safe_filename(contract.title)}-counsel-handoff.docx"
