"""PDF, DOCX, and text extraction with citation and quality metadata."""
from pathlib import Path


def _line_numbered(text):
    return "\n".join(f"[L{index}] {line}" for index, line in enumerate(text.splitlines(), start=1))


def _document_quality(full_text, documents, file_type, ocr_used=False):
    characters = len(full_text.strip())
    pages = len(documents) if file_type == "pdf" else 1
    warnings = []
    if characters < 250:
        warnings.append("Very little text was extracted. The document may be scanned or image-based.")
    if "\ufffd" in full_text:
        warnings.append("Some characters could not be decoded correctly.")
    return {
        "characters": characters,
        "pages": pages,
        "file_type": file_type.upper(),
        "ocr_used": ocr_used,
        "quality": "Needs review" if warnings else "Readable",
        "warnings": warnings,
    }


def parse_document(file_path):
    """Return line/page-marked text, citation-ready chunks, and extraction quality."""
    from langchain_core.documents import Document
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        from langchain_community.document_loaders import PyPDFLoader

        documents = PyPDFLoader(file_path).load()
        ocr_used = False
        if len("".join(doc.page_content for doc in documents).strip()) < 250:
            try:
                import pytesseract
                from pdf2image import convert_from_path

                images = convert_from_path(file_path, dpi=200)
                ocr_documents = []
                for index, image in enumerate(images, start=1):
                    ocr_documents.append(
                        Document(
                            page_content=pytesseract.image_to_string(image),
                            metadata={"page_label": index, "location": f"Page {index} (OCR)"},
                        )
                    )
                if len("".join(doc.page_content for doc in ocr_documents).strip()) > 250:
                    documents = ocr_documents
                    ocr_used = True
            except Exception:
                ocr_used = False
        marked_pages = []
        for index, doc in enumerate(documents, start=1):
            doc.metadata["page_label"] = index
            doc.metadata["location"] = f"Page {index}"
            marked_pages.append(f"[PAGE {index}]\n{doc.page_content}")
        full_text = "\n\n".join(marked_pages)
        file_type = "pdf"
    elif suffix == ".docx":
        from docx import Document as DocxDocument

        parsed = DocxDocument(file_path)
        paragraphs = [paragraph.text for paragraph in parsed.paragraphs if paragraph.text.strip()]
        raw_text = "\n".join(paragraphs)
        full_text = _line_numbered(raw_text)
        documents = [Document(page_content=raw_text, metadata={"location": "DOCX paragraphs"})]
        file_type = "docx"
        ocr_used = False
    else:
        raw_text = Path(file_path).read_text(encoding="utf-8", errors="replace")
        full_text = _line_numbered(raw_text)
        documents = [Document(page_content=raw_text, metadata={"location": "Text lines"})]
        file_type = "txt"
        ocr_used = False

    if not full_text.strip():
        return "", [], _document_quality("", documents, file_type, ocr_used)

    splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=140)
    chunks = splitter.split_documents(documents)
    for index, chunk in enumerate(chunks, start=1):
        chunk.metadata["chunk_id"] = index
        if not chunk.metadata.get("location"):
            chunk.metadata["location"] = f"Excerpt {index}"
    return full_text, chunks, _document_quality(full_text, documents, file_type, ocr_used)
